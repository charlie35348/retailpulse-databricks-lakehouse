
from collections import Counter
from datetime import datetime
from pathlib import Path
import hashlib
import json
import sqlite3
import tempfile

from openpyxl import load_workbook



SOURCE_FILE = Path.home() / "Placement" / "online_retail_II.xlsx"
OUTPUT_FILE = Path("reports") / "source_profile.json"

if not SOURCE_FILE.exists():
    raise FileNotFoundError(f"Dataset not found: {SOURCE_FILE}")

EXPECTED_COLUMNS = [
    "Invoice",
    "StockCode",
    "Description",
    "Quantity",
    "InvoiceDate",
    "Price",
    "Customer ID",
    "Country",
]

missing = Counter()
sheet_rows = Counter()
quantity_counts = Counter()
price_counts = Counter()
distinct = {
    "invoices": set(),
    "products": set(),
    "customers": set(),
    "countries": set(),
}

min_date = None
max_date = None
total_rows = 0
cancellation_rows = 0
blank_rows = 0

# SQLite stores only row fingerprints, keeping Python memory usage lower.
with tempfile.TemporaryDirectory(prefix="retailpulse_") as temp_dir:
    db_path = Path(temp_dir) / "row_fingerprints.sqlite"
    connection = sqlite3.connect(db_path)
    connection.execute(
        "CREATE TABLE seen_rows (fingerprint TEXT PRIMARY KEY)"
    )

    workbook = load_workbook(
        SOURCE_FILE,
        read_only=True,
        data_only=True,
    )

    try:
        actual_headers = None

        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            headers = next(rows, None)

            if headers is None:
                continue

            headers = [str(value).strip() for value in headers]

            if actual_headers is None:
                actual_headers = headers
                if headers != EXPECTED_COLUMNS:
                    raise ValueError(
                        f"Unexpected columns: {headers}\n"
                        f"Expected: {EXPECTED_COLUMNS}"
                    )
            elif headers != actual_headers:
                raise ValueError(
                    f"Inconsistent headers in {worksheet.title}: {headers}"
                )

            column_index = {
                name: index for index, name in enumerate(headers)
            }

            batch = []

            for values in rows:
                if all(
                    value is None
                    or (isinstance(value, str) and not value.strip())
                    for value in values
                ):
                    blank_rows += 1
                    continue

                total_rows += 1
                sheet_rows[worksheet.title] += 1

                record = dict(zip(headers, values))

                # Count nulls and whitespace-only strings.
                for column, value in record.items():
                    if value is None or (
                        isinstance(value, str) and not value.strip()
                    ):
                        missing[column] += 1

                invoice = record["Invoice"]
                stock = record["StockCode"]
                customer = record["Customer ID"]
                country = record["Country"]
                quantity = record["Quantity"]
                price = record["Price"]
                invoice_date = record["InvoiceDate"]

                if invoice is not None:
                    invoice_text = str(invoice).strip()
                    distinct["invoices"].add(invoice_text)

                    if invoice_text.upper().startswith("C"):
                        cancellation_rows += 1

                if stock is not None and str(stock).strip():
                    distinct["products"].add(str(stock).strip())

                if customer is not None and str(customer).strip():
                    distinct["customers"].add(str(customer).strip())

                if country is not None and str(country).strip():
                    distinct["countries"].add(str(country).strip())

                if isinstance(quantity, (int, float)):
                    if quantity < 0:
                        quantity_counts["negative"] += 1
                    elif quantity == 0:
                        quantity_counts["zero"] += 1
                    else:
                        quantity_counts["positive"] += 1

                if isinstance(price, (int, float)):
                    if price < 0:
                        price_counts["negative"] += 1
                    elif price == 0:
                        price_counts["zero"] += 1
                    else:
                        price_counts["positive"] += 1

                if isinstance(invoice_date, datetime):
                    if min_date is None or invoice_date < min_date:
                        min_date = invoice_date
                    if max_date is None or invoice_date > max_date:
                        max_date = invoice_date

                # Hash the entire original row for exact duplicate detection.
                canonical_row = json.dumps(
                    list(values),
                    ensure_ascii=False,
                    default=str,
                    separators=(",", ":"),
                )
                fingerprint = hashlib.sha256(
                    canonical_row.encode("utf-8")
                ).hexdigest()
                batch.append((fingerprint,))

                if len(batch) >= 2000:
                    connection.executemany(
                        "INSERT OR IGNORE INTO seen_rows VALUES (?)",
                        batch,
                    )
                    connection.commit()
                    batch.clear()

            if batch:
                connection.executemany(
                    "INSERT OR IGNORE INTO seen_rows VALUES (?)",
                    batch,
                )
                connection.commit()

    finally:
        workbook.close()

    unique_rows = connection.execute(
        "SELECT COUNT(*) FROM seen_rows"
    ).fetchone()[0]

    connection.close()

report = {
    "source_file": SOURCE_FILE.name,
    "source_size_mb": round(
        SOURCE_FILE.stat().st_size / (1024 * 1024), 2
    ),
    "worksheets": dict(sheet_rows),
    "total_transaction_rows": total_rows,
    "blank_rows_skipped": blank_rows,
    "exact_duplicate_rows": total_rows - unique_rows,
    "missing_values_by_column": {
        column: missing[column] for column in EXPECTED_COLUMNS
    },
    "quantity_counts": dict(quantity_counts),
    "price_counts": dict(price_counts),
    "cancellation_style_invoice_rows": cancellation_rows,
    "distinct_nonblank_counts": {
        name: len(values) for name, values in distinct.items()
    },
    "invoice_date_min": min_date.isoformat() if min_date else None,
    "invoice_date_max": max_date.isoformat() if max_date else None,
}

OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
OUTPUT_FILE.write_text(json.dumps(report, indent=2), encoding="utf-8")

print(json.dumps(report, indent=2))
print(f"\nSaved report to: {OUTPUT_FILE.resolve()}")
