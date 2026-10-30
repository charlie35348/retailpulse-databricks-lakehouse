
import csv
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook

SOURCE_FILE = Path.home() / "Placement" / "online_retail_II.xlsx"
OUTPUT_DIR = Path("data") / "raw"

EXPECTED_HEADERS = [
    "Invoice",
    "StockCode",
    "Description",
    "Quantity",
    "InvoiceDate",
    "Price",
    "Customer ID",
    "Country",
]


def format_value(value):
    """Convert Excel cell values to CSV-safe representations."""
    if value is None:
        return ""

    if isinstance(value, (datetime, date, time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()

    return str(value)


def main():
    if not SOURCE_FILE.exists():
        raise FileNotFoundError(
            f"Source workbook not found: {SOURCE_FILE}"
        )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(
        SOURCE_FILE,
        read_only=True,
        data_only=True,
    )

    try:
        total_exported = 0

        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            headers = next(rows, None)

            if headers is None:
                continue

            headers = [str(value).strip() for value in headers]

            if headers != EXPECTED_HEADERS:
                raise ValueError(
                    f"Unexpected headers in {worksheet.title}: {headers}"
                )

            filename = (
                "online_retail_2009_2010.csv"
                if worksheet.title == "Year 2009-2010"
                else "online_retail_2010_2011.csv"
            )
            output_file = OUTPUT_DIR / filename
            row_count = 0

            with output_file.open(
                "w",
                newline="",
                encoding="utf-8-sig",
            ) as csv_file:
                writer = csv.writer(csv_file)
                writer.writerow(headers)

                for row in rows:
                    # Skip only entirely blank rows.
                    if all(value is None for value in row):
                        continue

                    writer.writerow(
                        [format_value(value) for value in row]
                    )
                    row_count += 1

            total_exported += row_count

            print(
                f"{worksheet.title}: {row_count:,} rows -> "
                f"{output_file}"
            )

        print(f"\nTotal exported rows: {total_exported:,}")
        print(f"Output directory: {OUTPUT_DIR.resolve()}")

    finally:
        workbook.close()


if __name__ == "__main__":
    main()
