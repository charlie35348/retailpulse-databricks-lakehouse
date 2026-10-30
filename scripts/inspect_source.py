
from pathlib import Path
from openpyxl import load_workbook

# Change this to the actual location of your downloaded workbook.
FILE_PATH = Path.home() / "Placement" / "online_retail_II.xlsx"

if not FILE_PATH.exists():
    raise FileNotFoundError(
        f"Workbook not found: {FILE_PATH}\n"
        "Update FILE_PATH to match your downloaded filename."
    )

print(f"File: {FILE_PATH}")
print(f"Size: {FILE_PATH.stat().st_size / (1024 * 1024):.2f} MB")

# Read-only mode reduces memory usage for large Excel workbooks.
workbook = load_workbook(
    filename=FILE_PATH,
    read_only=True,
    data_only=True,
)

print("\nWorksheets:", workbook.sheetnames)

for worksheet in workbook.worksheets:
    print(f"\n--- Worksheet: {worksheet.title} ---")
    print(f"Reported rows: {worksheet.max_row}")
    print(f"Reported columns: {worksheet.max_column}")

    rows = worksheet.iter_rows(values_only=True)

    headers = next(rows, None)
    print("Headers:", headers)

    print("\nFirst 5 data rows:")
    for _, row in zip(range(5), rows):
        print(row)

workbook.close()
